"""Autonova D price-list parser.

The supplier did not provide a sample file yet, so the parser is deliberately
header-driven: it finds the header row, maps common Ukrainian/Russian/English
column names and supports both wide (one column per warehouse) and long
(one row per warehouse + quantity column) price layouts.
"""
from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import load_workbook


ARTICLE_ALIASES = [
    "артикул", "арт", "article", "art", "код товару", "код товара",
    "каталожний номер", "каталожный номер", "номер деталі", "номер детали",
    "номер запчастини", "номер запчастини", "part number", "partnumber",
    "p/n", "pn", "sku", "код", "номер",
]
BRAND_ALIASES = [
    "бренд", "виробник", "производитель", "brand", "producer", "manufacturer",
    "торгова марка", "торговая марка", "trade mark", "тм",
]
NAME_ALIASES = [
    "назва", "найменування", "наименование", "название", "name", "опис",
    "описание", "description", "товар", "part name", "деталь", "номенклатура",
]
RETAIL_ALIASES = [
    "ррц", "роздріб", "роздрібна", "розничная", "розница", "роздрібн",
    "розничн", "retail", "retail price", "рекомендована ціна",
    "рекомендованная цена",
]
WHOLESALE_ALIASES = [
    "опт", "оптова", "оптовая", "wholesale", "wholesale price",
    "закупівельна", "закупочная", "закупка", "закупки", "закупівля",
    "приход", "прихода", "прихід", "вхідна", "дилерська", "дилерская",
    "ціна", "цена", "price", "вартість", "стоимость",
]
CURRENCY_ALIASES = ["валюта", "currency", "currency code"]
CATEGORY_ALIASES = ["категорія", "категория", "category", "група", "группа", "group"]
IMAGE_ALIASES = ["зображення", "изображение", "image", "photo", "picture", "картинка"]
WAREHOUSE_ALIASES = ["склад", "складі", "складе", "філія", "филиал", "warehouse"]
STOCK_ALIASES = [
    "залишок", "остаток", "залишки", "остатки", "кількість", "количество",
    "кол", "кол-во", "кол во", "к-во", "к во", "наявність", "наявності",
    "наличие", "наличии", "на складі", "на складе", "stock", "qty",
    "quantity", "pcs", "шт",
]
TOTAL_STOCK_ALIASES = ["всього", "всего", "итого", "разом", "total"]


class AutonovaPriceParseError(ValueError):
    pass


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _identifier(value: Any) -> str:
    """Excel may store an article as a number; avoid trailing '.0'."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _norm(value: Any) -> str:
    text = _text(value).lower().replace("ё", "е").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def _key_norm(value: Any) -> str:
    text = _norm(value)
    text = re.sub(r"[^a-zа-яіїєґ0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _has_alias(header: Any, aliases: list[str]) -> bool:
    key = f" {_key_norm(header)} "
    for alias in aliases:
        alias_key = _key_norm(alias)
        if alias_key and f" {alias_key} " in key:
            return True
    return False


def _header_has_word(header: Any, words: tuple[str, ...]) -> bool:
    key = _key_norm(header)
    return any(word in key.split() for word in words)


def _classify_header(header: Any) -> str | None:
    if not _text(header):
        return None
    key = _key_norm(header)

    if any(word in key for word in ("упаков", "кратн", "pack")):
        return None

    # Price/stock columns are checked before "warehouse": headers like
    # "Ціна на складі" or "Залишок на складі" are not warehouse dimensions.
    if _has_alias(header, RETAIL_ALIASES):
        return "retail"
    if _has_alias(header, WHOLESALE_ALIASES):
        return "price"
    if _has_alias(header, STOCK_ALIASES) or _header_has_word(header, ("шт",)):
        return "stock"
    if any(word in key for word in ("склад", "warehouse", "філія", "филиал")):
        return "warehouse"
    if _has_alias(header, ARTICLE_ALIASES):
        return "article"
    if _has_alias(header, BRAND_ALIASES):
        return "brand"
    if _has_alias(header, NAME_ALIASES):
        return "name"
    if _has_alias(header, CURRENCY_ALIASES):
        return "currency"
    if _has_alias(header, CATEGORY_ALIASES):
        return "category"
    if _has_alias(header, IMAGE_ALIASES):
        return "image"
    if _has_alias(header, TOTAL_STOCK_ALIASES):
        return "stock_total"
    return None


def _classify_columns(headers: list[Any]) -> list[tuple[int, str, str | None]]:
    return [
        (idx, _text(header), _classify_header(header))
        for idx, header in enumerate(headers)
    ]


def _column_priority(header: Any, col_type: str) -> int:
    """Prefer the most specific column when several match the same type."""
    if col_type == "article":
        priorities = [
            ("артикул", 100),
            ("article", 95),
            ("каталож", 90),
            ("part number", 88),
            ("p n", 86),
            ("sku", 84),
            ("код товару", 80),
            ("код товара", 80),
            ("номер деталі", 75),
            ("номер детали", 75),
            ("номер запчастини", 74),
            ("номер", 60),
            ("код", 50),
        ]
        for alias, score in priorities:
            if _has_alias(header, [alias]):
                return score
    elif col_type == "price":
        if _has_alias(
            header,
            [
                "опт", "оптова", "оптовая", "wholesale", "дилерськ",
                "дилерск",
            ],
        ):
            return 100
        if _has_alias(
            header,
            [
                "закупівельн", "закупочн", "закупка", "закупки",
                "закупівля", "приход", "прихода", "прихід", "вхідна",
            ],
        ):
            return 90
        if _has_alias(header, ["ціна", "цена", "price"]):
            return 50
    elif col_type == "retail":
        if _has_alias(header, ["ррц"]):
            return 100
        return 50
    return 10


def _find_header(rows: list[tuple]) -> tuple[int, list[tuple[int, str, str | None]]]:
    best_idx = -1
    best_score = -1
    best_columns: list[tuple[int, str, str | None]] = []

    for idx, row in enumerate(rows[:50]):
        columns = _classify_columns(list(row))
        types = {col_type for _, _, col_type in columns if col_type}
        headers = [header for _, header, _ in columns if header]
        if not headers:
            continue

        score = 0
        if "article" in types:
            score += 3
        if "price" in types or "retail" in types:
            score += 2
        if "brand" in types:
            score += 2
        if "name" in types:
            score += 1
        if {"stock", "stock_total", "warehouse"} & types:
            score += 1

        if score > best_score:
            best_idx = idx
            best_score = score
            best_columns = columns

    if best_idx < 0 or best_score < 5 or "article" not in {
        col_type for _, _, col_type in best_columns if col_type
    }:
        preview = ", ".join(
            f"{i}:{header}" for i, header, _ in _classify_columns(list(rows[0])) if header
        ) if rows else ""
        raise AutonovaPriceParseError(
            "Autonova D: не удалось распознать строку заголовков прайса. "
            + (f"Колонки: {preview}" if preview else "Файл пустой.")
        )

    return best_idx, best_columns


def _parse_decimal(value: Any) -> float | None:
    if value is None:
        return None
    text = _text(value).replace("\xa0", "").replace(" ", "")
    if not text:
        return None
    text = text.replace("₴", "").replace("грн", "").replace("UAH", "")
    text = text.replace("$", "").replace("€", "")
    text = re.sub(r"[^0-9.,\-]", "", text)
    if not text:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        text = "".join(parts[:-1]) + "." + parts[-1] if len(parts) > 2 else text.replace(",", ".")

    try:
        return round(float(text), 2)
    except ValueError:
        return None


def _parse_stock(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return max(0, int(value))

    text = _text(value).lower()
    if not text:
        return 0
    match = re.search(r"-?\d+(?:[.,]\d+)?", text)
    if not match:
        if any(marker in text for marker in ("є", "есть", "+", "наяв", "налич")):
            return 1
        return 0
    try:
        return max(0, int(float(match.group().replace(",", "."))))
    except ValueError:
        return 0


def _parse_currency(value: Any) -> str:
    text = _text(value).upper()
    if "$" in text or "USD" in text:
        return "USD"
    if "€" in text or "EUR" in text:
        return "EUR"
    return "UAH"


def _build_warehouse_lookup(warehouses: list[dict] | None) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for warehouse in warehouses or []:
        name = _text(warehouse.get("Name") or warehouse.get("name"))
        if not name:
            continue
        for raw_key in (
            warehouse.get("Id"), warehouse.get("id"),
            warehouse.get("Abbreviation"), warehouse.get("abbreviation"),
            name,
        ):
            if raw_key is None:
                continue
            lookup[_norm(raw_key)] = name
            lookup[_key_norm(raw_key)] = name
            key_text = _text(raw_key)
            if key_text.isdigit():
                lookup[str(int(key_text))] = name
    return lookup


WAREHOUSE_STOP_WORDS = re.compile(
    r"(?i)\b(склад|складі|складе|філія|филиал|warehouse|залишок|остаток|"
    r"залишки|остатки|кількість|количество|кол-во|кол во|наявність|наличие|"
    r"stock|qty|quantity|шт|на|складі|складе)\b"
)


def _warehouse_name(value: Any, lookup: dict[str, str]) -> str:
    raw = _text(value)
    if not raw:
        return ""
    candidates = [_norm(raw), _key_norm(raw)]
    if raw.isdigit():
        candidates.append(str(int(raw)))
    for candidate in candidates:
        if candidate in lookup:
            return lookup[candidate]

    cleaned = re.sub(r"\s+", " ", WAREHOUSE_STOP_WORDS.sub(" ", raw)).strip(" ,.-")
    for candidate in (_norm(cleaned), _key_norm(cleaned)):
        if candidate and candidate in lookup:
            return lookup[candidate]

    return cleaned or raw


def _stock_column_key(header: str, lookup: dict[str, str]) -> str:
    cleaned = re.sub(
        r"\s+", " ", WAREHOUSE_STOP_WORDS.sub(" ", _text(header))
    ).strip(" ,.-")
    if not cleaned:
        return "stock"
    return _warehouse_name(cleaned, lookup) or cleaned


def _header_preview(columns: list[tuple[int, str, str | None]], limit: int = 30) -> str:
    return ", ".join(
        f"{idx}:{header}" for idx, header, _ in columns[:limit] if header
    )


def parse_autonova_price(
    file_data: bytes,
    warehouses: list[dict] | None = None,
) -> list[dict]:
    """Parse an Autonova D XLSX price list into SupplierPrice-shaped dicts."""
    try:
        workbook = load_workbook(
            io.BytesIO(file_data), read_only=True, data_only=True
        )
    except Exception as exc:
        raise AutonovaPriceParseError(
            f"Autonova D: не удалось открыть XLSX-прайс ({exc})"
        ) from exc

    warehouse_lookup = _build_warehouse_lookup(warehouses)
    best: tuple[int, int, list[tuple[int, str, str | None]], list[tuple]] | None = None
    fallback_preview = ""

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        if not fallback_preview:
            for row in rows[:10]:
                headers = [_text(value) for value in row if _text(value)]
                if headers:
                    fallback_preview = ", ".join(headers[:30])
                    break
        try:
            header_idx, columns = _find_header(rows)
        except AutonovaPriceParseError:
            continue
        score = len({col_type for _, _, col_type in columns if col_type})
        if best is None or score > best[0]:
            best = (score, header_idx, columns, rows)

    workbook.close()

    if best is None:
        raise AutonovaPriceParseError(
            "Autonova D: в XLSX-прайсе не найдена таблица с артикулами и ценами. "
            + (f"Колонки: {fallback_preview}" if fallback_preview else "")
        )

    _, header_idx, columns, rows = best
    mapping: dict[str, int] = {}
    priorities: dict[str, int] = {}
    stock_columns: list[tuple[int, str]] = []
    total_column: int | None = None
    warehouse_column: int | None = None

    for idx, header, col_type in columns:
        if col_type == "warehouse":
            warehouse_column = idx if warehouse_column is None else warehouse_column
        elif col_type == "stock_total":
            total_column = idx if total_column is None else total_column
        elif col_type == "stock":
            stock_columns.append((idx, header))
        elif col_type:
            priority = _column_priority(header, col_type)
            if col_type not in mapping or priority > priorities.get(col_type, -1):
                mapping[col_type] = idx
                priorities[col_type] = priority
        elif header:
            # A city/warehouse name in the header is a wide stock column if it
            # matches the warehouse directory from GET /api/v1/warehouses.
            normalized = _norm(header)
            key = _key_norm(header)
            if normalized in warehouse_lookup or key in warehouse_lookup:
                stock_columns.append((idx, header))

    # Fallback for price files whose warehouse columns are not in the
    # warehouse directory: numeric columns with neutral headers are stocks.
    if not stock_columns and warehouse_column is None and total_column is None:
        used_indexes = set(mapping.values())
        excluded_words = {
            "вага", "вес", "weight", "ндс", "пдв", "vat", "рік", "year",
            "гарант", "строк", "термін", "срок", "delivery", "днів", "дней",
            "days", "упаков", "кратн", "pack", "min", "max", "ціна", "цена",
            "price", "код", "code", "артикул", "article", "бренд", "brand",
            "виробник", "производитель", "назва", "наименование", "название",
            "name", "знижка", "скидка", "discount", "відсоток", "процент",
            "percent", "курс", "rate", "об", "єм", "объем", "volume",
            "розмір", "размер", "size", "діаметр", "диаметр", "дата", "date",
            "час", "время", "time", "оновлення", "обновление", "update",
            "lastupdate",
        }
        for idx, header, col_type in columns:
            if col_type is not None or not header or idx in used_indexes:
                continue
            key = _key_norm(header)
            if any(word in key.split() for word in excluded_words):
                continue
            values = [
                row[idx]
                for row in rows[header_idx + 1:]
                if any(value is not None and _text(value) != "" for value in row)
                and idx < len(row)
            ]
            if len(values) < 3:
                continue
            numeric = sum(
                1
                for value in values
                if (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                )
                or re.fullmatch(
                    r"\s*[<>]?\s*-?\d+(?:[.,]\d+)?\s*(?:шт\.?|од\.?|pcs?)?\s*",
                    _text(value),
                    re.IGNORECASE,
                )
            )
            if numeric / len(values) >= 0.8:
                stock_columns.append((idx, header))

    article_idx = mapping.get("article")
    if article_idx is None:
        raise AutonovaPriceParseError(
            "Autonova D: в прайсе не найдена колонка артикула. Колонки: "
            + _header_preview(columns)
        )

    brand_idx = mapping.get("brand")
    name_idx = mapping.get("name")
    price_idx = mapping.get("price")
    retail_idx = mapping.get("retail")
    currency_idx = mapping.get("currency")
    category_idx = mapping.get("category")
    image_idx = mapping.get("image")

    long_mode = warehouse_column is not None and bool(stock_columns)
    aggregated: dict[tuple[str, str], dict] = {}

    def cell(row: tuple, idx: int | None) -> Any:
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    for row in rows[header_idx + 1:]:
        if all(value is None or _text(value) == "" for value in row):
            continue

        article = _identifier(cell(row, article_idx))
        if not article or article.lower() in (
            "article", "артикул", "итого", "всего", "всього", "total"
        ):
            continue

        brand = _text(cell(row, brand_idx))
        name = _text(cell(row, name_idx))
        retail_price = _parse_decimal(cell(row, retail_idx))
        price = _parse_decimal(cell(row, price_idx))
        if price is None:
            price = retail_price
        currency = _parse_currency(cell(row, currency_idx))
        category = _text(cell(row, category_idx)) or None
        image_url = _text(cell(row, image_idx)) or None

        regions: dict[str, int] = {}
        stock_total = 0
        if long_mode:
            warehouse_raw = cell(row, warehouse_column)
            warehouse_name = _warehouse_name(warehouse_raw, warehouse_lookup)
            qty = _parse_stock(cell(row, stock_columns[0][0]))
            if warehouse_name:
                regions[warehouse_name] = regions.get(warehouse_name, 0) + qty
                stock_total += qty
        else:
            for idx, header in stock_columns:
                key = _stock_column_key(header, warehouse_lookup)
                qty = _parse_stock(cell(row, idx))
                regions[key] = regions.get(key, 0) + qty
                stock_total += qty
            if total_column is not None:
                explicit_total = _parse_stock(cell(row, total_column))
                if explicit_total > 0:
                    stock_total = explicit_total

        if not regions and stock_total > 0:
            regions["stock"] = stock_total

        key = (article, brand)
        if key in aggregated:
            existing = aggregated[key]
            # Long layout repeats one item per warehouse: sum stock. Wide
            # layout repeats the full stock picture: keep the last row.
            if long_mode:
                for warehouse_key, qty in regions.items():
                    existing["stock_regions"][warehouse_key] = (
                        existing["stock_regions"].get(warehouse_key, 0) + qty
                    )
                existing["stock_total"] += stock_total
            else:
                if regions:
                    existing["stock_regions"] = regions
                    existing["stock_total"] = stock_total
            if price is not None:
                existing["price"] = price
            if retail_price is not None:
                existing["retail_price"] = retail_price
            if name:
                existing["name"] = name
            if category:
                existing["category"] = category
            if image_url:
                existing["image_url"] = image_url
        else:
            aggregated[key] = {
                "article": article,
                "brand": brand,
                "name": name or article,
                "price": price,
                "retail_price": retail_price,
                "currency": currency,
                "stock_total": stock_total,
                "stock_regions": regions or None,
                "category": category,
                "image_url": image_url,
            }

    if not aggregated:
        raise AutonovaPriceParseError(
            "Autonova D: в прайсе не найдено ни одной строки с артикулом. "
            "Колонки: " + _header_preview(columns)
        )

    return list(aggregated.values())
